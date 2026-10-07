import os
import uuid
from pathlib import Path
from datetime import datetime
from docx import Document
from services.session import duration_label


def export_docx(session, text, target):
    if not text.strip():
        raise ValueError('No hay transcripción para exportar.')
    target = Path(target)
    document = Document()
    document.add_heading('Transcripción de consulta', 0)
    moment = datetime.fromisoformat(session.started_at.replace('Z', '+00:00'))
    for label, value in [('Código de paciente', session.patient_code or 'Sin código'),
                         ('Fecha', moment.strftime('%d/%m/%Y')),
                         ('Hora', moment.strftime('%H:%M:%S %z')),
                         ('Duración', duration_label(session.duration_seconds))]:
        document.add_paragraph(f'{label}: {value}')
    document.add_heading('Transcripción', 1)
    for paragraph in text.split('\n'):
        document.add_paragraph(paragraph)
    temporary = target.with_name(target.name + '.' + uuid.uuid4().hex + '.tmp')
    try:
        document.save(temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
